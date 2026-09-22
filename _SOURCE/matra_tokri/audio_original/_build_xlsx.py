"""Build AUDIO_SCRIPT.xlsx for मात्रा टोकरी (Matra Tokri).

Three columns:
  1. File Name        (no extension — the game uses <name>.mp3 automatically)
  2. Audio Text       (Hindi narration the VO artist should record)
  3. Format           ("mp3" — default, what the game expects)

This script lists every audio file the game actually fetches:
  - 15 word VOs (one per correct word caught: 5 per round x 3 rounds)
  -  1 tutorial line (vo-tutorial)
  -  3 round prompts (vo-round-aa / vo-round-i / vo-round-ee)
  -  3 praise lines between rounds (vo-good-1/2/3)
  -  1 win VO
  -  1 supplied FX file (sfx-incorrect — kept for consistency; see note)

ALL of the VO is already rendered with Gemini TTS (voice: Leda) — see
_build_vo.py to re-cut it, change voice, or hand this sheet to a human VO
artist and drop real recordings in over the top. The file names are the
contract; nothing in index.html changes.

NOTHING IS WRITTEN ON SCREEN except the falling words themselves, so every
one of these clips is load-bearing.

Catching a WRONG word intentionally plays no voice — the basket turns red and
wiggles and the word is tossed back out. Positive reinforcement only, exactly
as अक्षर वर्षा does with consonants.

The correct / wrong / round-complete / celebration SFX are the FLN Animation
Kit's CC0 recordings (audio/correct.ogg, wrong.ogg, burst.ogg,
sfx_celebrate.ogg) and need no recording. sfx-incorrect.mp3 is the team's own
wrong cue, kept in the folder as a one-line swap — see README.md.

Run:  python _build_xlsx.py
Output: AUDIO_SCRIPT.xlsx in the same folder.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROUNDS = [
    ("आ की मात्रा (Round 1) — plays when the word is caught", [
        ("s-naav",  "नाव"),
        ("s-haath", "हाथ"),
        ("s-baal",  "बाल"),
        ("s-kaam",  "काम"),
        ("s-daal",  "दाल"),
    ]),
    ("इ की मात्रा (Round 2) — plays when the word is caught", [
        ("s-din",   "दिन"),
        ("s-til",   "तिल"),
        ("s-sir",   "सिर"),
        ("s-dil",   "दिल"),
        ("s-hiran", "हिरन"),
    ]),
    ("बड़ी ई की मात्रा (Round 3) — plays when the word is caught", [
        ("s-nadi",    "नदी"),
        ("s-teer",    "तीर"),
        ("s-chini",   "चीनी"),
        ("s-machhli", "मछली"),
        ("s-neem",    "नीम"),
    ]),
]

VOS = [
    ("vo-tutorial", "टोकरी को उँगली से इधर-उधर ले जाओ।"),
    ("vo-round-aa", "आ की मात्रा वाले शब्दों को टोकरी में डालो।"),
    ("vo-round-i",  "अब इ की मात्रा वाले शब्दों को टोकरी में डालो।"),
    ("vo-round-ee", "अब बड़ी ई की मात्रा वाले शब्दों को टोकरी में डालो।"),
    ("vo-good-1",   "बहुत बढ़िया!"),
    ("vo-good-2",   "शाबाश!"),
    ("vo-good-3",   "कमाल कर दिया!"),
    ("vo-win",      "शाबाश! तुमने सभी मात्राओं के सही शब्दों को टोकरी में रख लिया है।"),
]

SFX = [
    ("sfx-incorrect", "Wrong-answer pop (already supplied by team — swiftee_incorrect.mp3). "
                      "Currently unused: the game plays the kit's wrong.ogg instead."),
]

wb = Workbook()
ws = wb.active
ws.title = "Audio Script"

HEADER_FILL  = PatternFill("solid", fgColor="002F76")
HEADER_FONT  = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="FCB717")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="002F76")
SOUND_FILL   = PatternFill("solid", fgColor="F0E6F8")
VO_FILL      = PatternFill("solid", fgColor="FFF5DA")
SFX_FILL     = PatternFill("solid", fgColor="E8F0E5")
CELL_FONT    = Font(name="Calibri", size=11, color="002F76")
HINDI_FONT   = Font(name="Nirmala UI", size=12, color="002F76")
THIN         = Side(border_style="thin", color="C8CFDA")
BORDER       = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER       = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT         = Alignment(horizontal="left",   vertical="center", wrap_text=True)


def header_row(row_idx, labels):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row_idx, column=col, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER


def section_row(row_idx, title, span=3):
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=span)
    c = ws.cell(row=row_idx, column=1, value=title)
    c.fill = SECTION_FILL
    c.font = SECTION_FONT
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.border = BORDER


def data_row(row_idx, file_name, audio_text, fill):
    cells = [
        (file_name, CELL_FONT, LEFT),
        (audio_text, HINDI_FONT, LEFT),
        ("mp3", CELL_FONT, CENTER),
    ]
    for col, (val, font, align) in enumerate(cells, start=1):
        c = ws.cell(row=row_idx, column=col, value=val)
        c.fill = fill
        c.font = font
        c.alignment = align
        c.border = BORDER


header_row(1, ["File Name", "Audio Text (Hindi)", "Format"])
ws.row_dimensions[1].height = 28

row = 2
total = 0

for title, words in ROUNDS:
    section_row(row, title)
    ws.row_dimensions[row].height = 22
    row += 1
    for name, text in words:
        data_row(row, name, text, SOUND_FILL)
        row += 1
        total += 1

section_row(row, "संदेश (Voice-overs) — tutorial, round prompts, praise, win")
ws.row_dimensions[row].height = 22
row += 1
for name, text in VOS:
    data_row(row, name, text, VO_FILL)
    row += 1
    total += 1

section_row(row, "SFX (already supplied — no re-recording needed)")
ws.row_dimensions[row].height = 22
row += 1
for name, text in SFX:
    data_row(row, name, text, SFX_FILL)
    row += 1
    total += 1

ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 74
ws.column_dimensions["C"].width = 12
ws.freeze_panes = "A2"

import os
folder = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(folder, "AUDIO_SCRIPT.xlsx")
try:
    wb.save(out)
except PermissionError:
    alt = os.path.join(folder, "AUDIO_SCRIPT_new.xlsx")
    wb.save(alt)
    out = alt
    print("NOTE: AUDIO_SCRIPT.xlsx was open/locked. Saved as AUDIO_SCRIPT_new.xlsx instead.")
print(f"Wrote {out}")
print(f"Total entries: {total}")
print(f"  Word VOs:     {sum(len(w) for _, w in ROUNDS)}")
print(f"  Voice-overs:  {len(VOS)}")
print(f"  SFX:          {len(SFX)}")
